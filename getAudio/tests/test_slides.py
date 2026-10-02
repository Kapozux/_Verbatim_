"""幻灯片（slides.py + /studio/slides、/studio/<id>/pptx、/studio/<id>/revise 接口）。

跑生产代码：check_slide / check（版式白名单、引语逐字核对、数字核对、出处换真 id）、fit_pt、render（真 python-pptx，
存完再用 python-pptx 打开检查：字是文本框不是图片、备注里有出处、页脚有期名）、slides.start（后台线程）、revise、接口。
桩（只在边界）：ask._llm（按提示词里真实出现的段落 id 造大纲：每种版式一页，外加编造的出处、对不上原文的引语、
原文里没有的数字、不认识的版式、不在词表里的图标、没写标题的结尾页）；slides.render_previews（要调 LibreOffice，太慢，
换成写一张假图，预览接口照样测）。
数据：一个合成博主，两期，每期几张卡（不碰资料库）。
"""
import json
import os
import re
import time

from _support import Checks, isolate, make_creator

TMP = isolate('slides')
import app as A  # noqa: E402
import ask  # noqa: E402
import slides  # noqa: E402
from pptx import Presentation  # noqa: E402

ok = {}
EP0 = ['Small teams ship faster than big ones.', 'We grew revenue 340 percent in two years.',
       'Talk to users every single week.']
EP1 = ['Hiring too early kills most startups.', 'In 2019 we had three people; by 2021 we had forty.']
CID = make_creator(A, 'Host', 'host', [EP0, EP1])
cdir = A._chain_dir(CID)
corpus = ask.load(cdir, passages=True)
by_q = {c['quote']: c['id'] for c in corpus['cards']}

prompts = []


def fake_llm(prompt, model=None, purpose='ask', grounded=False):
    prompts.append((purpose, prompt))
    i = lambda q: by_q[q]        # noqa: E731
    if 'Task: revise one slide' in prompt:
        return json.dumps({'slide': {'layout': 'points', 'title': 'Revised: small wins',
                                     'bullets': [f'Ship small and often [#{i(EP0[0])}]', 'Made up [#9-99]'],
                                     'notes': f'Because [#{i(EP0[0])}]'}})
    deck = [
        {'layout': 'cover', 'title': 'How small teams win', 'subtitle': 'What the host keeps saying', 'notes': ''},
        {'layout': 'points', 'title': 'Small beats big',
         'bullets': [{'text': f'Small teams ship faster [#{i(EP0[0])}]', 'detail': 'Fewer handoffs', 'icon': 'rocket'},
                     {'text': f'Weekly user talks [#{i(EP0[2])}]', 'icon': 'not-an-icon'}, 'Invented [#7-77]'],
         'notes': f'Small teams win because they move faster. [#{i(EP0[0])}]'},
        {'layout': 'quote', 'title': 'In his words', 'quote': EP0[0], 'who': 'Host', 'cite': f'[#{i(EP0[0])}]', 'notes': ''},
        {'layout': 'quote', 'title': 'Reworded quote', 'quote': 'Little teams are quicker than large ones.',
         'who': 'Host', 'cite': i(EP0[0]), 'notes': ''},                                  # 对不上原文 → 换成卡片原话
        {'layout': 'stat', 'title': 'Growth', 'value': '340%', 'label': 'revenue growth in two years',
         'cite': f'[#{i(EP0[1])}]', 'notes': ''},
        {'layout': 'stat', 'title': 'Fake number', 'value': '999%', 'label': 'nonsense', 'cite': f'[#{i(EP0[1])}]'},  # 原文没有 → 丢
        {'layout': 'compare', 'title': 'Early vs late hiring',
         'left': {'label': 'Too early', 'bullets': [f'Kills startups [#{i(EP1[0])}]']},
         'right': {'label': 'Small team', 'bullets': [f'Ships faster [#{i(EP0[0])}]']}, 'notes': ''},
        {'layout': 'timeline', 'title': 'From three to forty',
         'steps': [{'when': '2019', 'what': f'Three people [#{i(EP1[1])}]'}, {'when': '2021', 'what': f'Forty people [#{i(EP1[1])}]'}],
         'notes': ''},
        {'layout': 'hologram', 'title': 'Unknown layout'},                                   # 不认识、没要点 → 丢
        {'layout': 'section', 'title': 'Part two', 'subtitle': '', 'notes': ''},
        {'layout': 'closing', 'title': '', 'bullets': [{'text': f'Stay small [#{i(EP0[0])}]', 'icon': 'check'}], 'notes': ''},
    ]
    return '```json\n' + json.dumps({'slides': deck}) + '\n```'


ask._llm = fake_llm


def fake_previews(pptx, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    with open(os.path.join(out_dir, '1.png'), 'wb') as f:
        f.write(b'\x89PNG fake')
    return 1


slides.render_previews = fake_previews

# 1 字号自动缩：长文字在小框里会缩，短的保持最大
ok['fit'] = slides.fit_pt(['短'], 10, 1, 28) == 28 and slides.fit_pt(['很长的一句话' * 20], 6, 1, 28) < 20

# 2 生成一份（同步）
o = slides.start(cdir, CID, 'detailed', 6, run=lambda f: f())
ss = (o.get('result') or {}).get('slides') or []
lays = [s['layout'] for s in ss]
print('2 生成：', o['status'], o.get('error'), lays, '| 丢出处', o.get('dropped_citations'), '| 费用', o.get('cost_usd'))
ok['generate'] = o['status'] == 'done' and o.get('title') == 'How small teams win' and o.get('n_slides') == len(ss)
ok['layouts'] = lays == ['cover', 'points', 'quote', 'quote', 'stat', 'compare', 'timeline', 'section', 'closing']
ok['quote_verbatim'] = ss[3]['quote'] == EP0[0] and ss[3]['cite'] == [by_q[EP0[0]]]
ok['stat'] = ss[4]['value'] == '340%' and not any(s.get('value') == '999%' for s in ss)
b0, b1 = ss[1]['bullets'][0], ss[1]['bullets'][1]
ok['icons'] = (b0['icon'] == 'rocket' and b0['detail'] == 'Fewer handoffs' and b1['icon'] == ''
               and ss[-1]['title'] == 'What to remember' and 'Pick one name from this list' not in ''
               and 'rocket' in next(p for purpose, p in prompts if purpose == 'slides'))
ok['cites'] = (o['dropped_citations'] >= 1 and all(c in o['citations'] for s in ss for c in slides._slide_ids(s))
               and '7-77' not in json.dumps(ss) and not any('[#' in b['text'] for s in ss for b in s.get('bullets') or []))
script_prompt = next(p for purpose, p in prompts if purpose == 'slides')
ok['prompt'] = 'Task: plan a slide deck of 6 content slides' in script_prompt and 'Style: detailed' in script_prompt

# 3 文件：python-pptx 打开，字是文本框（能改）、备注有出处、页脚有期名
path = slides.file_path(cdir, o['id'])
prs = Presentation(path)
texts = [[sh.text_frame.text for sh in sl.shapes if sh.has_text_frame] for sl in prs.slides]
pics = sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13)
notes = [sl.notes_slide.notes_text_frame.text for sl in prs.slides]
print('3 文件：', len(prs.slides), '页 |', texts[1][:3], '| 备注', notes[1][:80].replace('\n', ' / '))
ok['pptx'] = (len(prs.slides) == len(ss) and pics > 0 and abs(prs.slide_width.inches - 13.333) < 0.01
              and all(sh.image.content_type == 'image/png' for sl in prs.slides for sh in sl.shapes if sh.shape_type == 13)
              and 'Fewer handoffs' in ' '.join(texts[1]) and o.get('previews') == 1
              and 'How small teams win' in texts[0] and 'Small teams ship faster' in ' '.join(texts[1])
              and any(t.startswith('Sources: ') or t.startswith('出处：') for t in texts[1])
              and EP0[0] in notes[1] and 'EP1' in ' '.join(texts[1]))

# 3b 排版细节：章节按章节序号编号（不是页码）、只有一条要点的页、结尾页带说明、长数字不折行
extra = [{'layout': 'cover', 'title': 'T', 'subtitle': '', 'notes': '', 'icon': ''},
         {'layout': 'section', 'title': 'Part A', 'subtitle': '', 'notes': '', 'icon': ''},
         {'layout': 'points', 'title': 'One thing', 'notes': 'The single point matters most.', 'icon': '',
          'bullets': [{'text': 'Only one bullet', 'detail': 'With its detail', 'icon': 'idea', 'cite': []}]},
         {'layout': 'section', 'title': 'Part B', 'subtitle': '', 'notes': '', 'icon': ''},
         {'layout': 'stat', 'title': 'Big', 'value': '340.3 万亿元', 'label': 'M2', 'cite': [], 'notes': '', 'icon': ''},
         {'layout': 'closing', 'title': 'End', 'notes': '', 'icon': '',
          'bullets': [{'text': 'Keep this', 'detail': 'Because of that', 'icon': 'check', 'cite': []},
                      {'text': 'And this', 'detail': 'For this reason', 'icon': 'star', 'cite': []}]}]
xp = os.path.join(TMP, 'extra.pptx')
slides.render(extra, xp, {}, 'T', False)
xt = [[sh.text_frame.text for sh in sl.shapes if sh.has_text_frame] for sl in Presentation(xp).slides]
ok['layout_details'] = ('02' in xt[3] and '04' not in xt[3] and 'Only one bullet' in xt[2] and 'With its detail' in xt[2]
                        and 'Because of that' in ' '.join(xt[5]) and '340.3 万亿元' in xt[4])

# 4 改一页：只换那一页，编造的出处删掉，文件重排
o2 = slides.revise(cdir, CID, o['id'], 1, 'make it shorter')
ss2 = o2['result']['slides']
prs2 = Presentation(path)
t2 = ' '.join(sh.text_frame.text for sh in prs2.slides[1].shapes if sh.has_text_frame)
ok['revise'] = (ss2[1]['title'] == 'Revised: small wins' and ss2[2] == ss[2] and len(ss2) == len(ss)
                and [b['text'] for b in ss2[1]['bullets']] == ['Ship small and often', 'Made up']
                and ss2[1]['bullets'][1]['cite'] == [] and 'Revised: small wins' in t2 and o2.get('revised_at'))
try:
    slides.revise(cdir, CID, o['id'], 99, 'x')
    ok['revise_bad'] = False
except ValueError:
    ok['revise_bad'] = True

# 5 接口：估价、做一份（后台）、列表、下载、改一页、删掉连 pptx 一起删
c = A.app.test_client()
ok['estimate'] = 0 < c.get(f'/api/chain/{CID}/studio/slides/estimate').get_json()['usd'] < 1
oid = c.post(f'/api/chain/{CID}/studio/slides', json={'style': 'presenter', 'n': 10}).get_json()['item']['id']
for _ in range(100):
    it = next(x for x in c.get(f'/api/chain/{CID}/studio').get_json()['items'] if x['id'] == oid)
    if it['status'] != 'running':
        break
    time.sleep(0.1)
dl = c.get(f'/api/chain/{CID}/studio/{oid}/pptx')
png = c.get(f'/api/chain/{CID}/studio/{oid}/slide/1')
png404 = c.get(f'/api/chain/{CID}/studio/{oid}/slide/9')
rv = c.post(f'/api/chain/{CID}/studio/{oid}/revise', json={'i': 2, 'instruction': 'shorter'})
rv_bad = c.post(f'/api/chain/{CID}/studio/{oid}/revise', json={'i': 2, 'instruction': ''})
c.delete(f'/api/chain/{CID}/studio/{oid}')
print('5 接口：', it['status'], it.get('title'), dl.status_code, dl.mimetype, dl.headers.get('Content-Disposition'), rv.status_code, rv_bad.status_code)
ok['api'] = (it['status'] == 'done' and it.get('kind') == 'slides' and it.get('style') == 'presenter'
             and dl.status_code == 200 and dl.data[:2] == b'PK' and 'How small teams win.pptx' in (dl.headers.get('Content-Disposition') or '')
             and rv.status_code == 200 and rv.get_json()['result']['slides'][2]['title'] == 'Revised: small wins'
             and png.status_code == 200 and png.mimetype == 'image/png' and png404.status_code == 404
             and rv_bad.status_code == 400 and not os.path.exists(slides.file_path(cdir, oid))
             and not os.path.exists(slides.preview_dir(cdir, oid))
             and c.get(f'/api/chain/{CID}/studio/{oid}/pptx').status_code == 404)

t = Checks()
for k, v in ok.items():
    t.check(k, v)
t.finish()
